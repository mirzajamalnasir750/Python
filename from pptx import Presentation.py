from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_PARAGRAPH_ALIGNMENT

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Palette
NAVY = RGBColor(18, 34, 62)
BLUE = RGBColor(43, 104, 190)
TEAL = RGBColor(0, 153, 153)
PALE = RGBColor(239, 245, 252)
WHITE = RGBColor(255,255,255)
DARK = RGBColor(38, 50, 68)
MUTED = RGBColor(91, 107, 128)
ORANGE = RGBColor(242, 153, 74)
GREEN = RGBColor(49, 145, 99)
RED = RGBColor(196, 76, 76)

def set_bg(slide, color=WHITE):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color

def add_text(slide, x, y, w, h, text, size=20, color=DARK, bold=False,
             font="Aptos", align=None, valign=MSO_ANCHOR.TOP, margin=0.05):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return shape

def add_rich_bullet(slide, x, y, w, h, items, size=20, color=DARK, gap=12):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.level = 0
        p.text = "•  " + item
    return shape

def add_header(slide, title, subtitle=None, num=None):
    # accent line
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.18)).fill.solid()
    bar = slide.shapes[-1]
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    add_text(slide, 0.65, 0.42, 12, 0.6, title, 28, NAVY, True)
    if subtitle:
        add_text(slide, 0.68, 1.08, 12, 0.42, subtitle, 13, MUTED)
    if num is not None:
        add_text(slide, 12.3, 7.08, 0.45, 0.22, f"{num}/7", 10, MUTED, align=PP_ALIGN.RIGHT)

def add_card(slide, x,y,w,h,title,body,accent=BLUE, body_size=18):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = PALE
    sh.line.color.rgb = RGBColor(220, 229, 240)
    # accent stripe
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.10), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    add_text(slide, x+0.25,y+0.18,w-0.45,0.45,title,20,NAVY,True)
    add_rich_bullet(slide,x+0.25,y+0.72,w-0.45,h-0.82,body,body_size,gap=8)

# Slide 1
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s,NAVY)
# Decorative circles
for x,y,d,c in [(10.7,0.7,1.3,TEAL),(11.7,1.5,0.8,BLUE),(10.1,2.0,0.55,ORANGE)]:
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.fill.background()
add_text(s,0.9,1.15,8.8,0.45,"BASIC OF AI",16,RGBColor(128,220,218),True)
add_text(s,0.9,2.0,9.5,1.35,"AI & MACHINE LEARNING",34,WHITE,True)
add_text(s,0.92,3.35,8.8,0.85,"From the big picture to Classification vs Regression",22,RGBColor(219,231,245))
add_text(s,0.95,5.35,8.9,0.75,"Mirza Muhammad Jamal Nasir",17,WHITE)
add_text(s,0.95,6.55,9.5,0.3,"Learning goal: identify the difference between classification and regression.",12,RGBColor(191,205,224))
add_text(s,12.1,7.08,0.5,0.22,"1/7",10,RGBColor(191,205,224),align=PP_ALIGN.RIGHT)

# Slide 2
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s,"1. What is Artificial Intelligence?","AI is the big picture: systems performing tasks that usually require human-like intelligence.",2)
add_card(s,0.7,1.75,5.85,4.75,"AI in everyday life",[
    "Voice assistants ",
    "Maps ",
    "Recommended App's",
    "Email services help detect spam"
],TEAL,19)
add_card(s,6.85,1.75,5.75,4.75,"Key idea",[
    "AI is a broad field of computing",
    "It includes different methods and technologies",
    "Machine Learning is one important approach within AI",
    ""
],BLUE,19)
add_text(s,0.85,6.65,11.7,0.35,"Remember: AI is the wider field like Class Machine Learning is one part of it like sub Class.",15,TEAL,True)

# Slide 3
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s,"2. AI vs Machine Learning","Think of AI as the wider goal, and ML as one way to achieve it.",3)
# Nested-like conceptual cards
add_card(s,0.8,1.75,5.75,4.6,"ARTIFICIAL INTELLIGENCE (AI)",[
    "The broader field",
    "Builds systems that perform intelligent tasks",
    "May use rules, search, planning, or learning"
],TEAL,19)
add_card(s,6.8,1.75,5.75,4.6,"MACHINE LEARNING (ML)",[
    "A subfield of AI",
    "Learns patterns from examples and data",
    "Uses learned patterns to predict or make decisions"
],BLUE,19)
add_text(s,1.05,6.55,11.1,0.45,"Example: learn from labelled spam emails instead of writing a rule for every possible message.",16,NAVY,True)

# Slide 4
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s,"3. Four common Machine Learning approaches","Different approaches depend on the feedback or labels available during learning.",4)
cards=[
    (0.65,1.75,"SUPERVISED","Learns from labelled examples.\nExample: spam / not spam.",TEAL),
    (6.85,1.75,"UNSUPERVISED","Finds patterns in unlabelled data.\nExample: grouping customers.",BLUE),
    (0.65,4.1,"SEMI-SUPERVISED","Uses a small labelled set plus more unlabelled data.",ORANGE),
    (6.85,4.1,"REINFORCEMENT","Learns actions using rewards or penalties.\nExample: an agent learning a game.",GREEN)
]
for x,y,t,b,c in cards:
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.8), Inches(1.95))
    sh.fill.solid(); sh.fill.fore_color.rgb=PALE; sh.line.color.rgb=RGBColor(220,229,240)
    sh2=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x+0.22), Inches(y+0.25), Inches(0.35), Inches(0.35))
    sh2.fill.solid(); sh2.fill.fore_color.rgb=c; sh2.line.fill.background()
    add_text(s,x+0.72,y+0.18,4.8,0.42,t,17,NAVY,True)
    add_text(s,x+0.25,y+0.78,5.25,0.95,b,16,DARK)
add_text(s,0.8,6.45,11.9,0.4,"Today's focus: Supervised Learning → Classification and Regression",16,TEAL,True)

# Slide 5
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s,"4. Supervised Learning: two prediction tasks","Training examples include inputs and known target answers (labels).",5)
add_card(s,0.7,1.8,5.85,4.7,"CLASSIFICATION",[
    "Predicts a category or class",
    "Spam or not spam",
    "Cat or dog",
    "Pass or fail"
],TEAL,20)
add_card(s,6.8,1.8,5.85,4.7,"REGRESSION",[
    "Predicts a numerical value",
    "House price: £250,000",
    "Temperature: 21.5°C",
    "Monthly sales: £15,000"
],BLUE,20)
add_text(s,1.0,6.65,11.3,0.35,"Quick rule: category → classification | numerical amount → regression",16,NAVY,True)

# Slide 6
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s,"5. How to identify the right task","Ask: what kind of answer are we trying to predict?",6)
add_card(s,0.75,1.8,5.75,3.95,"Ask: “Which category?”",[
    "Spam or not spam?",
    "Fraud or legitimate?",
    "Customer will leave or stay?",
    "These are classification examples."
],TEAL,18)
add_card(s,6.8,1.8,5.75,3.95,"Ask: “How much / how many?”",[
    "What will the house cost?",
    "What will tomorrow's temperature be?",
    "How much sales revenue is expected?",
    "These are regression examples."
],BLUE,18)
add_text(s,0.95,6.15,11.4,0.65,"Note: Classification models may output probabilities, but the target is a class. Regression predicts a numerical quantity.",14,MUTED)

# Slide 7
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s,"6. Quick check & takeaway","Pause for a moment and choose your answer before checking.",7)
add_card(s,0.7,1.65,5.85,2.05,"Question 1",[
    "Predict whether an email is spam.",
    "Classification or Regression?"
],TEAL,18)
add_card(s,6.8,1.65,5.85,2.05,"Question 2",[
    "Predict the selling price of a house.",
    "Classification or Regression?"
],BLUE,18)
add_card(s,0.7,4.0,5.85,1.65,"Answer 1",[
    "Classification — the output is a category."
],GREEN,17)
add_card(s,6.8,4.0,5.85,1.65,"Answer 2",[
    "Regression — the output is a number."
],GREEN,17)
add_text(s,0.9,6.15,11.7,0.7,"Takeaway: AI is the broad field • ML learns from data • Supervised Learning includes classification and regression tasks.",15,NAVY,True)

# Set presentation metadata
prs.core_properties.title = "AI and Machine Learning: Classification vs Regression"
prs.core_properties.subject = "Five-minute lecturer micro-teach for Computing"
prs.core_properties.author = "OpenAI"
prs.core_properties.keywords = "Artificial Intelligence, Machine Learning, Supervised Learning, Classification, Regression"

out = r"C:\Users\Zain\OneDrive\Desktop\mirzajamalnasir750\AI_Machine_Learning_Classification_vs_Regression_Arden_Microteach.pptx"
prs.save(out)
print(out)


